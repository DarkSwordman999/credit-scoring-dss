"""Выгрузка таблиц в XLSX (С-16): заголовок, параметры, таблица с автошириной."""
import io
import re
from datetime import date, datetime
from urllib.parse import quote

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="999999")
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")           # недопустимы в имени листа Excel
_BAD_FILE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')   # недопустимы в имени файла Windows


def build(title: str, headers: list[str], rows: list[list], meta: list[tuple[str, str]] | None = None,
          sheet: str = "Отчёт", footer: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _BAD_SHEET.sub("-", sheet)[:31] or "Отчёт"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in meta or []:
        ws.append([k, v])
    ws.append([])
    head_row = ws.max_row + 1
    ws.append(headers)
    for c in ws[head_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3EAF4")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    for r in rows:
        ws.append([_cell(v) for v in r])
    for row in ws.iter_rows(min_row=head_row + 1, max_row=ws.max_row):
        for c in row:
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
            elif isinstance(c.value, (date, datetime)):
                c.number_format = "DD.MM.YYYY HH:MM" if isinstance(c.value, datetime) else "DD.MM.YYYY"
    for line in footer or []:
        ws.append([])
        ws.append([line])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(_cell(r[i - 1]) or "")) for r in rows[:500] if len(r) >= i])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if isinstance(v, float):
        return round(v, 2)
    return v


def response(data: bytes, filename: str) -> Response:
    filename = _BAD_FILE.sub("-", filename)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


def build_tables(title: str, meta: list[tuple[str, str]], summary: list[tuple[str, object]],
                 tables: list[dict]) -> bytes:
    """Отчёт из нескольких таблиц: первый лист — параметры и итоги, далее лист на каждую таблицу."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Итоги"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in list(meta) + [("Сформирован", datetime.now().strftime("%d.%m.%Y %H:%M"))] + list(summary):
        ws.append([k, _cell(v)])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30
    used = {"Итоги"}
    for t in tables:
        name = _BAD_SHEET.sub("-", t["title"])[:31] or "Таблица"
        while name in used:
            name = name[:28] + f"-{len(used)}"
        used.add(name)
        sh = wb.create_sheet(name)
        sh.append(t["headers"])
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E3EAF4")
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        for r in t["rows"]:
            sh.append([_cell(v) for v in r])
        for row in sh.iter_rows(min_row=2, max_row=sh.max_row):
            for c in row:
                c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
                if isinstance(c.value, datetime):
                    c.number_format = "DD.MM.YYYY HH:MM"
                elif isinstance(c.value, date):
                    c.number_format = "DD.MM.YYYY"
        for i, h in enumerate(t["headers"], start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in t["rows"][:300] if len(r) >= i])
            sh.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
