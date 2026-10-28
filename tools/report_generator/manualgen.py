# -*- coding: utf-8 -*-
"""Эксплуатационные документы (руководство пользователя, руководство по эксплуатации) в docx и PDF.

Оформление — по ГОСТ 19.106-78 / ГОСТ 2.105-2019 и РД 50-34.698-90: титульный лист с полями согласования
и утверждения, аннотация, содержание (поле TOC, обновляется в MS Word), разделы с нумерацией,
рисунки и таблицы с подписями, лист регистрации изменений. Нумерация страниц — внизу по центру,
титульный лист не нумеруется."""
import os

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import docgen


class Numbers:
    """Сквозная нумерация рисунков и таблиц: fig() и table() возвращают блоки docgen."""

    def __init__(self):
        self.f = 0
        self.t = 0

    def fig(self, path, caption, width=16.5, max_h=21.0):
        self.f += 1
        return ("fig", path, f"Рисунок {self.f} – {caption}", width, max_h)

    def next_fig(self, k=1):
        return self.f + k

    def table(self, caption, header, rows, widths):
        self.t += 1
        return ("table", str(self.t), caption, header, rows, widths)

    def next_table(self, k=1):
        return self.t + k


def _line(doc, text="", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=0, upper=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.15
    pf.space_after = Pt(after)
    if text:
        docgen._add_rich(p, text.upper() if upper else text, size=size, bold=bold)
    return p


def _field(p, code, size=14):
    run = p.add_run()
    docgen._set_font(run, size)
    for tag, text in (("begin", None), (None, code), ("separate", None), (None, "—"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tag)
            run._r.append(fc)
        elif text == code:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {code} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def title(doc, info):
    """info: system, edition, doc, code, year, [org]."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Cm(docgen.TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    for k, (a, b) in enumerate((("СОГЛАСОВАНО", "УТВЕРЖДАЮ"), ("", ""),
                                ("________________ / ____________ /", "________________ / ____________ /"),
                                ("«___» ____________ " + info["year"] + " г.", "«___» ____________ " + info["year"] + " г."))):
        if k:
            p.add_run().add_break()
        r = p.add_run(f"{a}\t{b}")   # не _add_rich: «__» в полях подписи — не разметка курсива
        docgen._set_font(r, 14, bold=(k == 0))
    for _ in range(6):
        _line(doc)
    for t in info["system"].split("\n"):
        _line(doc, t, bold=True, upper=True)
    if info.get("edition"):
        _line(doc, info["edition"], after=18)
    _line(doc, info["doc"], bold=True, size=16, after=6, upper=True)
    _line(doc, info["code"], after=6)
    p = _line(doc, "Листов ")
    _field(p, "NUMPAGES")
    for _ in range(12):
        _line(doc)
    p = _line(doc, info["year"])
    p.add_run().add_break(WD_BREAK.PAGE)


def toc(doc):
    p = docgen.para(doc, "Содержание", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    for r in p.runs:
        r.font.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    _field(p, 'TOC \\o "1-2" \\h \\z \\u')
    # следующий раздел начнётся с новой страницы: заголовки разделов — с разрывом перед ними


def change_sheet(doc):
    """Лист регистрации изменений (ГОСТ 19.604-78, ГОСТ 2.503-2013)."""
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.page_break_before = True
    docgen._add_rich(p, "Лист регистрации изменений", bold=True)
    head1 = ["Изм.", "Номера листов (страниц)", "", "", "", "Всего листов (страниц) в докум.",
             "№ докум.", "Входящий № сопроводительного докум. и дата", "Подп.", "Дата"]
    head2 = ["", "изменённых", "заменённых", "новых", "аннулированных", "", "", "", "", ""]
    widths = [1.0, 1.9, 1.9, 1.5, 2.1, 1.7, 1.4, 2.4, 1.3, 1.3]
    n_rows = 22
    t = doc.add_table(rows=2 + n_rows, cols=10)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    docgen._set_col_widths(t, widths)
    for r, head in enumerate((head1, head2)):
        for c, txt in enumerate(head):
            docgen._cell_text(t.cell(r, c), txt, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    t.cell(0, 1).merge(t.cell(0, 4))
    docgen._cell_text(t.cell(0, 1), "Номера листов (страниц)", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c in (0, 5, 6, 7, 8, 9):
        m = t.cell(0, c).merge(t.cell(1, c))
        docgen._cell_text(m, head1[c], size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in range(2, 2 + n_rows):
        t.rows[r].height = Cm(0.8)


def build(blocks, out_path, info, annotation, splits=None):
    doc = docgen.new_document()
    # каждый раздел и приложение — с новой страницы (ГОСТ 19.106-78, п. 2.1.4)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    title(doc, info)
    docgen.render(doc, [("hcn", "Аннотация")] + [("p", a) for a in annotation] + [("pb",)])
    toc(doc)
    docgen.render(doc, blocks, splits)
    change_sheet(doc)
    doc.save(out_path)


def _update_fields(word, path, analyze=False):
    """Обновить содержание и поля (число листов); analyze — заодно найти, где разрезать таблицы."""
    d = word.Documents.Open(os.path.abspath(path))
    try:
        # строки содержания — одинарный интервал (стили «Оглавление 1, 2»: wdStyleTOC1 = -20, wdStyleTOC2 = -21),
        # иначе длинное содержание с хвостовым абзацем выталкивает пустую страницу
        for st in (-20, -21):
            pf = d.Styles(st).ParagraphFormat
            pf.LineSpacingRule = 0
            pf.SpaceBefore = 0
            pf.SpaceAfter = 0
        for i in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(i).Update()
        d.Fields.Update()
        d.Repaginate()
        for i in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(i).UpdatePageNumbers()
        d.Fields.Update()
        d.Save()
        return docgen.analyze_splits(d) if analyze else None
    finally:
        d.Close(False)


def build_final(blocks, out_path, info, annotation, word=None, pdf=True):
    """Сборка с разрезанием таблиц по страницам («Продолжение таблицы N»), обновлением содержания и PDF."""
    import win32com.client
    own = word is None
    if own:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
    splits = {}
    try:
        for _ in range(60):
            build(blocks, out_path, info, annotation, splits)
            changes = _update_fields(word, out_path, analyze=True)
            if not changes:
                break
            _, num, kind, idx = changes[0]
            lst = splits.setdefault(num, [])
            if kind == "add" and idx not in lst and idx > 0:
                lst.append(idx)
            elif kind == "del" and idx in lst:
                lst.remove(idx)
            else:
                break
        pages = None
        if pdf:
            pages = docgen.word_export_pdf(word, out_path, os.path.splitext(out_path)[0] + ".pdf")
        return splits, pages
    finally:
        if own:
            word.Quit()
